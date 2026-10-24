import os
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    # 16:9 widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Palette (Enterprise Supply Chain Palette)
    NAVY = RGBColor(15, 23, 42)          # Slate 900
    SLATE_DARK = RGBColor(30, 41, 59)     # Slate 800
    SLATE_LIGHT = RGBColor(241, 245, 249) # Slate 100
    WHITE = RGBColor(255, 255, 255)
    ACCENT_BLUE = RGBColor(2, 132, 199)   # Sky 600
    ACCENT_GREEN = RGBColor(22, 163, 74)  # Green 600
    TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500
    BORDER_COLOR = RGBColor(203, 213, 225)# Slate 300
    RED_ACCENT = RGBColor(185, 28, 28)    # Red 700

    def add_header(slide, title_text, category_text="STOCKWISE AI | TRIWIZATHON PS-2"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.5), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(23)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

    def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    tbox1 = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.5), Inches(3.6))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "TRIWIZATHON PROBLEM STATEMENT 2"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = ACCENT_BLUE
    p_badge.space_after = Pt(8)

    p_main = tf1.add_paragraph()
    p_main.text = "StockWise AI"
    p_main.font.size = Pt(46)
    p_main.font.bold = True
    p_main.font.color.rgb = WHITE
    p_main.space_after = Pt(14)

    p_sub = tf1.add_paragraph()
    p_sub.text = "AI-Powered FMCG Distribution Intelligence & Inventory Optimization Engine"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = RGBColor(226, 232, 240)
    p_sub.space_after = Pt(18)

    p_tag = tf1.add_paragraph()
    p_tag.text = "A local-first, mathematically sound decision engine connecting demand forecasting, multi-echelon replenishment, and constraint-aware Mixed-Integer Programming."
    p_tag.font.size = Pt(14)
    p_tag.font.color.rgb = RGBColor(148, 163, 184)

    fbox1 = slide1.shapes.add_textbox(Inches(1.2), Inches(5.8), Inches(10.5), Inches(0.8))
    p_foot = fbox1.text_frame.paragraphs[0]
    p_foot.text = "Architecture: Next.js 15 | FastAPI | Google OR-Tools | SQLite3 | 12-Invariant Mathematical Verifier"
    p_foot.font.size = Pt(11.5)
    p_foot.font.color.rgb = RGBColor(148, 163, 184)

    # =========================================================================
    # SLIDE 2: The Operational Problem
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "The FMCG Operational Problem: Multi-Variable Distribution Dilemma")

    cards2 = [
        ("Demand Volatility Across Corridors",
         "Consumer order rates swing dynamically across retail tiers (hypermarkets, supermarkets, mom-and-pop bodegas) with pronounced day-of-week seasonality and Poisson order clustering."),
        ("Uneven Multi-Echelon Stock",
         "Regional warehouses experience severe inventory imbalances: central hubs hold surplus while peripheral hubs risk imminent stockouts on high-velocity SKUs."),
        ("Perishable Expiry & FEFO Pressure",
         "Shipping older batches across slow transit routes risks arrival with insufficient shelf life, generating retail shrinkage and FMCG distributor return penalties."),
        ("Transport Cost vs. Service Level",
         "Expedited deliveries prevent store stockouts but erode FMCG operating margins. Planners must balance freight budgets against customer fulfillment guarantees.")
    ]

    for idx, (title, desc) in enumerate(cards2):
        col = idx % 2
        row = idx // 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.5 + row * 2.7)
        add_card(slide2, left, top, Inches(5.6), Inches(2.4), SLATE_LIGHT, BORDER_COLOR)
        
        tbox = slide2.shapes.add_textbox(left + Inches(0.3), top + Inches(0.25), Inches(5.0), Inches(1.9))
        tf = tbox.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = f"0{idx+1}. {title}"
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = NAVY
        p1.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = SLATE_DARK

    # =========================================================================
    # SLIDE 3: Our Solution
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "StockWise AI Solution: Closed-Loop Distribution Intelligence")

    steps = [
        ("1. Demand Forecasting", "EWMA, Naive & Seasonal-Naive over 800 corridors with backtesting validation (MAE & WAPE)."),
        ("2. Replenishment Calculation", "Lead-time demand + dynamic safety stock to compute exact carton replenishment thresholds."),
        ("3. Constraint-Aware MIP", "Google OR-Tools solver evaluating physical stock, throughput limits, and delivery cutoffs."),
        ("4. Independent Verification", "12-invariant external verifier guarantees feasibility before any plan is presented."),
        ("5. Explainable Execution", "Driver-based natural language attribution and RFC 4180 CSV export for supply chain planners.")
    ]

    for idx, (title, desc) in enumerate(steps):
        left = Inches(0.8 + idx * 2.35)
        top = Inches(1.7)
        add_card(slide3, left, top, Inches(2.2), Inches(5.0), WHITE, BORDER_COLOR)

        tbox = slide3.shapes.add_textbox(left + Inches(0.15), top + Inches(0.25), Inches(1.9), Inches(4.5))
        tf = tbox.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13.5)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_BLUE
        p1.space_after = Pt(12)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = SLATE_DARK

    # =========================================================================
    # SLIDE 4: System Architecture
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Local-First Enterprise System Architecture")

    layers = [
        ("Presentation Layer (Next.js 15 App Router)", "React 19, TypeScript strict mode, TanStack React Query v5, Recharts, Tailwind CSS tokens. 7 interactive operational control towers."),
        ("API Services Layer (FastAPI Asynchronous Backend)", "Pydantic v2 schemas, OpenAPI/Swagger contracts, RFC 4180 / OWASP CSV sanitization, planning run management."),
        ("Analytical Engines Layer (Forecasting & Replenishment)", "EWMA / Naive / Seasonal Naive models, lead-time demand scaling, dynamic safety stock, stockout risk banding."),
        ("Optimization & Verification Layer (OR-Tools MIP & Verifier)", "Google OR-Tools branch-and-cut linear solver, 3 multi-objective profiles, independent 12-invariant mathematical auditor."),
        ("Persistence Layer (Relational SQLite3 & Alembic)", "Master data, routes, batch-level inventory, 96k demand history records, persisted planning runs, and shipment audit trails.")
    ]

    for idx, (layer_title, layer_desc) in enumerate(layers):
        top = Inches(1.5 + idx * 1.1)
        add_card(slide4, Inches(0.8), top, Inches(7.5), Inches(0.95), SLATE_LIGHT, BORDER_COLOR)
        tbox = slide4.shapes.add_textbox(Inches(1.0), top + Inches(0.1), Inches(7.1), Inches(0.75))
        tf = tbox.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = layer_title
        p1.font.size = Pt(12.5)
        p1.font.bold = True
        p1.font.color.rgb = NAVY
        p2 = tf.add_paragraph()
        p2.text = layer_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = SLATE_DARK

    add_card(slide4, Inches(8.6), Inches(1.5), Inches(3.9), Inches(5.35), WHITE, BORDER_COLOR)
    rtbox = slide4.shapes.add_textbox(Inches(8.8), Inches(1.7), Inches(3.5), Inches(4.9))
    rtf = rtbox.text_frame
    rtf.word_wrap = True
    rp1 = rtf.paragraphs[0]
    rp1.text = "Key Architectural Pillars"
    rp1.font.size = Pt(16)
    rp1.font.bold = True
    rp1.font.color.rgb = NAVY
    rp1.space_after = Pt(12)

    pillars = [
        ("Strict Plan Decoupling", "Advisory plans in PlanningRun never mutate base inventory items in InventoryItem."),
        ("100% Deterministic Seed", "Seed 42 generates identical coordinates, routes, and batch expiries on every clean run."),
        ("Independent Greedy Baseline", "Built-in heuristic comparator evaluates identical constraints without solver bias."),
        ("Offline & Local-First", "Zero external cloud or LLM dependencies. Operates seamlessly in secure air-gapped environments.")
    ]
    for p_title, p_desc in pillars:
        pt = rtf.add_paragraph()
        pt.text = f"• {p_title}:"
        pt.font.size = Pt(11.5)
        pt.font.bold = True
        pt.font.color.rgb = ACCENT_BLUE
        pd = rtf.add_paragraph()
        pd.text = f"  {p_desc}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = SLATE_DARK
        pd.space_after = Pt(8)

    # =========================================================================
    # SLIDE 5: Demand & Inventory Intelligence
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Demand Forecasting & Inventory Accounting Intelligence")

    add_card(slide5, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.35), WHITE, BORDER_COLOR)
    ltbox = slide5.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.2), Inches(4.9))
    ltf = ltbox.text_frame
    ltf.word_wrap = True
    lp1 = ltf.paragraphs[0]
    lp1.text = "Forecasting vs. Replenishment Derivation"
    lp1.font.size = Pt(16)
    lp1.font.bold = True
    lp1.font.color.rgb = NAVY
    lp1.space_after = Pt(10)

    f_points = [
        ("Consumer Demand Forecasting", "EWMA (alpha=0.30) models daily mean demand across 800 corridors from 120-day empirical history. Validated via backtesting holdouts (MAE & WAPE)."),
        ("Lead-Time Demand (LTD)", "LTD = Daily Mean Demand x Route Lead Time in days across fastest eligible lane. Network sum: 28,745.44 units."),
        ("Dynamic Safety Stock (SS)", "SS = z x sigma x sqrt(L) under standard cycle-service level assumption (z=1.65, 95% cycle target; not guaranteed volume fill rate). Network sum: 10,744.93 units."),
        ("Replenishment Requirement", "Total Demand = ceil(LTD + SS) = 39,878 units. Represents integer cartons needed to prevent stockouts over transit lead times.")
    ]
    for pt, pd in f_points:
        p = ltf.add_paragraph()
        p.text = f"• {pt}: {pd}"
        p.font.size = Pt(11)
        p.font.color.rgb = SLATE_DARK
        p.space_after = Pt(9)

    add_card(slide5, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.35), SLATE_LIGHT, BORDER_COLOR)
    rtbox5 = slide5.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.3), Inches(4.9))
    rtf5 = rtbox5.text_frame
    rtf5.word_wrap = True
    rp1_5 = rtf5.paragraphs[0]
    rp1_5.text = "4-Tier Inventory Accounting & FEFO Policy"
    rp1_5.font.size = Pt(16)
    rp1_5.font.bold = True
    rp1_5.font.color.rgb = NAVY
    rp1_5.space_after = Pt(10)

    tiers = [
        ("1. Physical On-Hand Units", "Gross warehouse physical stock recorded in warehouse bins."),
        ("2. Reserved Units", "Stock already earmarked for existing orders; strictly unavailable for new allocations."),
        ("3. Available Units", "On-Hand minus Reserved units. Upper bound before quality and expiry checks."),
        ("4. Usable Units (FEFO Strict)", "Available stock filtered for quality: Quarantined batches (0 usable) and Expired batches (0 usable) are completely disqualified from allocation.")
    ]
    for tt, td in tiers:
        p = rtf5.add_paragraph()
        p.text = f"{tt}:"
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        pd = rtf5.add_paragraph()
        pd.text = f"{td}"
        pd.font.size = Pt(11)
        pd.font.color.rgb = SLATE_DARK
        pd.space_after = Pt(7)

    # =========================================================================
    # SLIDE 6: Optimization Engine
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Google OR-Tools MIP Optimization Engine: Formulations & Constraints")

    profiles = [
        ("Cost-First Profile",
         "Minimizes pure transportation freight subject to a minimum fulfillment floor.",
         "min sum(c_wd * x) + sum(2*P * u)\ns.t. sum(x) >= ceil(0.80 * Demand)",
         "2*P unmet penalty. Bypassed if unconstrained opt-in. Trims high-cost lanes while exceeding 80% floor (delivers 90.18%)."),
        ("Service-First Profile",
         "True Lexicographic 2-Stage Priority Solve: Maximizes fulfillment first, then minimizes cost.",
         "Stage 1: min sum(u) -> yields u*\nStage 2: min sum(c_wd * x) s.t. sum(u) <= u*",
         "Guarantees maximum customer service (91.74%) with globally optimal routing ($52,812.09)."),
        ("Balanced Profile",
         "Multi-objective formulation balancing freight, transit speed, and lost-sales penalties.",
         "min sum[(c_wd + 0.20*T_wd)*x] + sum(5*P * u)\nUnits: All terms in $/unit",
         "Balances freight cost against product freshness (freshness penalty: $0.20/day/unit).")
    ]

    for idx, (p_title, p_sub, p_math, p_desc) in enumerate(profiles):
        left = Inches(0.8 + idx * 3.9)
        top = Inches(1.5)
        add_card(slide6, left, top, Inches(3.7), Inches(4.3), WHITE, BORDER_COLOR)

        tbox = slide6.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(3.3), Inches(3.9))
        tf = tbox.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = p_title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = NAVY
        p1.space_after = Pt(3)

        p2 = tf.add_paragraph()
        p2.text = p_sub
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_after = Pt(8)

        p3 = tf.add_paragraph()
        p3.text = p_math
        p3.font.size = Pt(10.5)
        p3.font.bold = True
        p3.font.color.rgb = ACCENT_BLUE
        p3.space_after = Pt(8)

        p4 = tf.add_paragraph()
        p4.text = p_desc
        p4.font.size = Pt(10.5)
        p4.font.color.rgb = SLATE_DARK

    # Bottom Constraint Callout Box
    add_card(slide6, Inches(0.8), Inches(6.0), Inches(11.7), Inches(1.15), SLATE_LIGHT, BORDER_COLOR)
    cbox6 = slide6.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(1.05))
    ctf6 = cbox6.text_frame
    ctf6.word_wrap = True
    cp1_6 = ctf6.paragraphs[0]
    cp1_6.text = "Mathematical Objective vs. Reported Freight Metric & Hard Constraints"
    cp1_6.font.size = Pt(11.5)
    cp1_6.font.bold = True
    cp1_6.font.color.rgb = NAVY
    cp2_6 = ctf6.add_paragraph()
    cp2_6.text = "• Objective value includes internal mathematical penalties (2*P*u, 0.20*T*x); reported total_transport_cost is strictly actual freight: sum(x * c_wd).\n• Enforced Hard Constraints: Warehouse usable supply bounds | Route eligibility & delivery cutoffs | Warehouse dispatch caps (15k u/day) | Destination receiving caps (3k u/day) | Non-negative integers (x in Z+)."
    cp2_6.font.size = Pt(10)
    cp2_6.font.color.rgb = SLATE_DARK

    # =========================================================================
    # SLIDE 7: Measured Results & Baseline Comparison
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Empirical Benchmark Validation: Optimization vs. Greedy Baseline")

    rows = 7
    cols = 5
    table_shape = slide7.shapes.add_table(rows, cols, Inches(0.8), Inches(1.5), Inches(11.7), Inches(3.1))
    table = table_shape.table

    table_data = [
        ["Metric", "Cost-First", "Balanced", "Service-First", "Greedy Baseline"],
        ["Solver Status", "OPTIMAL", "OPTIMAL", "OPTIMAL", "HEURISTIC"],
        ["Total Demand Considered", "39,878 u", "39,878 u", "39,878 u", "39,878 u"],
        ["Fulfilled Units (Rate %)", "35,963 u (90.18%)", "36,583 u (91.74%)", "36,583 u (91.74%)", "36,583 u (91.74%)"],
        ["Total Transportation Cost", "$50,150.99", "$52,816.51", "$52,812.09", "$54,085.66"],
        ["Avg Cost / Fulfilled Unit", "$1.3945 / u", "$1.4437 / u", "$1.4436 / u", "$1.4784 / u"],
        ["Freight vs. Baseline", "-$3,934.67 (-7.27%)", "-$1,269.15 (-2.35%)", "-$1,273.57 (-2.35%)", "Baseline Reference"]
    ]

    for r_idx, row in enumerate(table_data):
        for c_idx, cell_value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = cell_value
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            if r_idx == 0:
                p.font.bold = True
                p.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            else:
                p.font.color.rgb = NAVY
                if r_idx % 2 == 1:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = SLATE_LIGHT

    add_card(slide7, Inches(0.8), Inches(4.85), Inches(11.7), Inches(2.3), WHITE, BORDER_COLOR)
    cbox = slide7.shapes.add_textbox(Inches(1.0), Inches(4.95), Inches(11.3), Inches(2.1))
    ctf = cbox.text_frame
    ctf.word_wrap = True

    cp1 = ctf.paragraphs[0]
    cp1.text = "Mathematical Analysis of Benchmark Findings (Deterministic Seed: 42)"
    cp1.font.size = Pt(13.5)
    cp1.font.bold = True
    cp1.font.color.rgb = ACCENT_BLUE
    cp1.space_after = Pt(4)

    cp2 = ctf.add_paragraph()
    cp2.text = "1. Fulfillment Parity (91.74%): Service-First, Balanced, and Baseline all reach exactly 36,583 fulfilled units—the network's physical supply ceiling (3,295 units unmet due to network-wide SKU-0004 butter shortage)."
    cp2.font.size = Pt(11)
    cp2.font.color.rgb = SLATE_DARK
    cp2.space_after = Pt(3)

    cp3 = ctf.add_paragraph()
    cp3.text = "2. Pure Routing Efficiency (-2.35%): Service-First fulfills the identical 36,583 units as the greedy heuristic for $52,812.09 vs. $54,085.66, saving $1,273.57 in pure freight through global multi-commodity allocation."
    cp3.font.size = Pt(11)
    cp3.font.color.rgb = SLATE_DARK
    cp3.space_after = Pt(3)

    cp4 = ctf.add_paragraph()
    cp4.text = "3. Cost-First Margin Protection (-7.27%): Cost-First saves $3,934.67 in freight by trimming 620 units (-1.56%) of margin-diluting long-haul deliveries while strictly exceeding the 80% fulfillment floor (delivering 90.18%)."
    cp4.font.size = Pt(11)
    cp4.font.color.rgb = SLATE_DARK

    # =========================================================================
    # SLIDE 8: Explainability & Verification
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Decision Explainability & Independent 12-Invariant Verification")

    add_card(slide8, Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.35), WHITE, BORDER_COLOR)
    vtbox = slide8.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.6), Inches(4.9))
    vtf = vtbox.text_frame
    vtf.word_wrap = True
    vp1 = vtf.paragraphs[0]
    vp1.text = "Independent 12-Invariant Solution Verifier"
    vp1.font.size = Pt(16)
    vp1.font.bold = True
    vp1.font.color.rgb = NAVY
    vp1.space_after = Pt(8)

    v_points = [
        "1. Strictly Positive Quantities (x > 0)",
        "2. Integer Allocations (x in Z+)",
        "3. Usable Batch Status (0 for expired/quarantined)",
        "4. Warehouse Supply Bounds (allocated <= usable)",
        "5. Demand Ceilings (no destination over-fulfillment)",
        "6. Active Route Validity (inactive routes rejected)",
        "7. Warehouse Dispatch Caps (<= 15,000 u/day)",
        "8. Destination Receiving Caps (<= 3,000 u/day)",
        "9. Delivery Lead-Time Cutoff Compliance",
        "10. Batch Arrival Shelf-Life Feasibility Check",
        "11. Demand Reconciliation (Fulfilled + Unmet == Demanded)",
        "12. Independent Total Cost Recalculation (sum(x * c) == Cost)"
    ]
    for vp in v_points:
        p = vtf.add_paragraph()
        p.text = f"✓ {vp}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = ACCENT_GREEN

    add_card(slide8, Inches(7.1), Inches(1.5), Inches(5.4), Inches(5.35), SLATE_LIGHT, BORDER_COLOR)
    etbox = slide8.shapes.add_textbox(Inches(7.3), Inches(1.7), Inches(5.0), Inches(4.9))
    etf = etbox.text_frame
    etf.word_wrap = True
    ep1 = etf.paragraphs[0]
    ep1.text = "Driver-Based Decision Attribution"
    ep1.font.size = Pt(16)
    ep1.font.bold = True
    ep1.font.color.rgb = NAVY
    ep1.space_after = Pt(8)

    expl_items = [
        ("No Hallucinated Rationale", "Explanations cite only exact numbers from the solver contract (stock volume, route lead time, freight cost)."),
        ("Source Attribution", "Identifies winning warehouse based on verified available inventory and lowest-cost eligible transit."),
        ("Rejection Justifications", "Explicitly details why competing warehouses were rejected (e.g. WH-04 rejected due to stock exhaustion, WH-02 rejected due to +$0.45/u freight penalty)."),
        ("Root-Cause Deficit Explanations", "When unmet demand occurs, attributes the deficit to physical supply depletion rather than route cutoffs.")
    ]
    for et, ed in expl_items:
        p = etf.add_paragraph()
        p.text = f"• {et}:"
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        pd = etf.add_paragraph()
        pd.text = f"  {ed}"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = SLATE_DARK
        pd.space_after = Pt(8)

    # =========================================================================
    # SLIDE 9: Live Demonstration Workflow
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "5-Minute Interactive Live Demonstration Workflow")

    demo_steps = [
        ("01. Overview Control Tower", "0:00 - 0:45", "Verify LOCAL-FIRST status, multi-echelon network scope (4 WH, 16 SKUs, 50 Destinations), and warehouse capacity telemetry."),
        ("02. Inventory Explorer", "0:45 - 1:30", "Audit 4-tier inventory ledger. Filter by QUARANTINED to prove that quarantined batches show 0 Usable units and are excluded."),
        ("03. Demand Forecasting", "1:30 - 2:15", "Select Corridor SKU-0001 (Whole Milk) -> DST-01 (Sacramento). Run EWMA (alpha=0.30) 7-day forecast. Display Recharts curve and backtest holdout evaluation (MAE/WAPE)."),
        ("04. Replenishment Planner", "2:15 - 3:00", "Review stockout urgency tiers. Highlight GAP-1 / GAP-2 operational notices. Open explanation drawer showing LTD and Safety Stock math."),
        ("05. Local OR-Tools Solver", "3:00 - 3:45", "Execute Balanced profile. Showcase sub-second solve (0.20s), OPTIMAL status, Verifier: Feasible badge, and shipment schedule."),
        ("06. Scenario Matrix & Export", "3:45 - 5:00", "Run 3-profile comparison matrix proving cost vs service trade-offs. Download RFC 4180 CSV export with formula injection sanitization.")
    ]

    for idx, (step_title, step_time, step_desc) in enumerate(demo_steps):
        col = idx % 3
        row = idx // 3
        left = Inches(0.8 + col * 3.9)
        top = Inches(1.5 + row * 2.7)
        add_card(slide9, left, top, Inches(3.7), Inches(2.4), WHITE, BORDER_COLOR)

        tbox = slide9.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(3.3), Inches(2.0))
        tf = tbox.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = step_title
        p1.font.size = Pt(13.5)
        p1.font.bold = True
        p1.font.color.rgb = NAVY

        p_time = tf.add_paragraph()
        p_time.text = f"Target Time: {step_time}"
        p_time.font.size = Pt(10.5)
        p_time.font.bold = True
        p_time.font.color.rgb = ACCENT_BLUE
        p_time.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = step_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = SLATE_DARK

    # =========================================================================
    # SLIDE 10: Limitations & Production Next Steps
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Architectural Disclosures & Production Roadmap")

    add_card(slide10, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.35), SLATE_LIGHT, BORDER_COLOR)
    ltbox10 = slide10.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(5.3), Inches(4.9))
    ltf10 = ltbox10.text_frame
    ltf10.word_wrap = True
    lp1_10 = ltf10.paragraphs[0]
    lp1_10.text = "Current Limitations & Disclosures"
    lp1_10.font.size = Pt(16)
    lp1_10.font.bold = True
    lp1_10.font.color.rgb = NAVY
    lp1_10.space_after = Pt(8)

    limits = [
        ("Simulated FMCG Dataset", "Synthetic network (seed 42) reflects realistic FMCG dynamics, but results represent simulated benchmarks rather than realized enterprise savings."),
        ("GAP-1: Destination Inventory Unmodeled", "Retail stores currently modeled with 0 initial inventory; demand represents gross replenishment needed over lead time."),
        ("GAP-2: In-Transit Pipeline Unmodeled", "Prior in-transit shipments not tracked in current schema; treated as 0 incoming units."),
        ("Store-Level Expiry Unmodeled", "Warehouse expiry strictly enforced (FEFO); store POS shelf spoilage unmodeled and honestly labeled 'Not Modeled'."),
        ("Advisory Planning", "Produces human decision support; does not trigger physical dispatches without planner approval.")
    ]
    for lt, ld in limits:
        p = ltf10.add_paragraph()
        p.text = f"• {lt}:"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = RED_ACCENT
        pd = ltf10.add_paragraph()
        pd.text = f"  {ld}"
        pd.font.size = Pt(10)
        pd.font.color.rgb = SLATE_DARK
        pd.space_after = Pt(4)

    add_card(slide10, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.35), WHITE, BORDER_COLOR)
    rtbox10 = slide10.shapes.add_textbox(Inches(7.0), Inches(1.7), Inches(5.3), Inches(4.9))
    rtf10 = rtbox10.text_frame
    rtf10.word_wrap = True
    rp1_10 = rtf10.paragraphs[0]
    rp1_10.text = "Production Next Steps (Roadmap)"
    rp1_10.font.size = Pt(16)
    rp1_10.font.bold = True
    rp1_10.font.color.rgb = NAVY
    rp1_10.space_after = Pt(8)

    roadmap = [
        ("1. Real-Time ERP & WMS Connectors", "Direct connectors to SAP / Oracle SCM to ingest live inventory lots, batch statuses, and confirmed dispatch streams."),
        ("2. Store-Level POS & Shrinkage Ingestion", "Close GAP-1 by tracking store backroom on-hand stock and actual register checkout velocities."),
        ("3. GPS Telematics Pipeline", "Close GAP-2 by streaming live in-transit truck positions and real-time highway delay re-routing."),
        ("4. Multi-Period Rolling Horizon MIP", "Extend static single-period optimization into a multi-echelon 14-day rolling horizon MILP with inventory carryover."),
        ("5. Automated WMS Dispatch Execution", "Trigger automated ASN (Advanced Shipping Notice) webhooks upon supply chain planner approval.")
    ]
    for rmt, rmd in roadmap:
        p = rtf10.add_paragraph()
        p.text = f"• {rmt}:"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        pd = rtf10.add_paragraph()
        pd.text = f"  {rmd}"
        pd.font.size = Pt(10)
        pd.font.color.rgb = SLATE_DARK
        pd.space_after = Pt(4)

    output_path = REPO_ROOT / "docs" / "StockWise_AI_Hackathon_Presentation.pptx"
    prs.save(str(output_path))
    print(f"[+] Successfully generated PowerPoint presentation: {output_path}")

if __name__ == "__main__":
    REPO_ROOT = Path(r"c:\Users\dell\Desktop\ruk-pr")
    build_presentation()
